from flask import Flask, render_template, request, redirect, url_for, flash, send_file
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from supabase import create_client
from config import Config
from datetime import date
import calendar
import uuid
from flask_login import LoginManager, UserMixin, login_user, logout_user, login_required, current_user
from werkzeug.security import check_password_hash, generate_password_hash

app = Flask(__name__)
app.config.from_object(Config)

supabase = create_client(app.config["SUPABASE_URL"], app.config["SUPABASE_KEY"])

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "login"


from functools import wraps

class User(UserMixin):
    def __init__(self, id, username, nama, role):
        self.id = id
        self.username = username
        self.nama = nama
        self.role = role


@login_manager.user_loader
def load_user(user_id):
    data = supabase.table("users").select("*").eq("id", user_id).single().execute().data
    if data:
        return User(data["id"], data["username"], data.get("nama"), data.get("role", "staff"))
    return None


def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if current_user.role != "admin":
            flash("Halaman ini khusus untuk admin.", "error")
            return redirect(url_for("dashboard"))
        return f(*args, **kwargs)
    return decorated


@app.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        username = request.form["username"]
        password = request.form["password"]
        data = supabase.table("users").select("*").eq("username", username).execute().data
        if data and check_password_hash(data[0]["password_hash"], password):
            user = User(data[0]["id"], data[0]["username"], data[0].get("nama"), data[0].get("role", "staff"))
            login_user(user)
            flash("Berhasil login.", "success")
            next_page = request.args.get("next")
            return redirect(next_page or url_for("dashboard"))
        flash("Username atau password salah.", "error")

    return render_template("login.html")


@app.route("/logout")
@login_required
def logout():
    logout_user()
    flash("Berhasil logout.", "success")
    return redirect(url_for("login"))


def get_bulan_terakhir(n=6):
    bulan_list = []
    today = date.today()
    y, m = today.year, today.month
    for i in range(n):
        bulan_list.append((y, m))
        m -= 1
        if m == 0:
            m = 12
            y -= 1
    bulan_list.reverse()
    return bulan_list


@app.route("/")
@login_required
def dashboard():
    transaksi = supabase.table("transaksi").select("*").order("tanggal", desc=True).limit(10).execute().data
    total_masuk = supabase.table("transaksi").select("jumlah").eq("jenis", "masuk").execute().data
    total_keluar = supabase.table("transaksi").select("jumlah").eq("jenis", "keluar").execute().data

    sum_masuk = sum(t["jumlah"] for t in total_masuk)
    sum_keluar = sum(t["jumlah"] for t in total_keluar)
    saldo = sum_masuk - sum_keluar

    bulan_list = get_bulan_terakhir(6)
    awal = date(bulan_list[0][0], bulan_list[0][1], 1)
    akhir_tahun, akhir_bulan = bulan_list[-1]
    akhir_hari = calendar.monthrange(akhir_tahun, akhir_bulan)[1]
    akhir = date(akhir_tahun, akhir_bulan, akhir_hari)

    data_chart = supabase.table("transaksi").select("tanggal, jenis, jumlah") \
        .gte("tanggal", str(awal)).lte("tanggal", str(akhir)).execute().data

    chart_labels = []
    chart_masuk = []
    chart_keluar = []
    nama_bulan = ["", "Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Ags", "Sep", "Okt", "Nov", "Des"]

    for (y, m) in bulan_list:
        chart_labels.append(f"{nama_bulan[m]} {y}")
        key = f"{y}-{m:02d}"
        total_m = sum(t["jumlah"] for t in data_chart if t["jenis"] == "masuk" and t["tanggal"][:7] == key)
        total_k = sum(t["jumlah"] for t in data_chart if t["jenis"] == "keluar" and t["tanggal"][:7] == key)
        chart_masuk.append(total_m)
        chart_keluar.append(total_k)

    return render_template("dashboard.html",
                           transaksi=transaksi,
                           sum_masuk=sum_masuk,
                           sum_keluar=sum_keluar,
                           saldo=saldo,
                           chart_labels=chart_labels,
                           chart_masuk=chart_masuk,
                           chart_keluar=chart_keluar)


def upload_bukti(file):
    if not file or file.filename == "":
        return None
    ext = file.filename.rsplit(".", 1)[-1].lower()
    nama_file = f"{uuid.uuid4()}.{ext}"
    file_bytes = file.read()
    supabase.storage.from_("bukti-transaksi").upload(
        nama_file, file_bytes, {"content-type": file.content_type}
    )
    url = supabase.storage.from_("bukti-transaksi").get_public_url(nama_file)
    url = url.rstrip("?")
    return url


@app.route("/transaksi")
@login_required
def list_transaksi():
    jenis = request.args.get("jenis")
    kategori_id = request.args.get("kategori_id")
    tanggal_mulai = request.args.get("tanggal_mulai")
    tanggal_akhir = request.args.get("tanggal_akhir")
    q = request.args.get("q", "").strip()

    query = supabase.table("transaksi").select("*, kategori(nama)")

    if jenis:
        query = query.eq("jenis", jenis)
    if kategori_id:
        query = query.eq("kategori_id", kategori_id)
    if tanggal_mulai:
        query = query.gte("tanggal", tanggal_mulai)
    if tanggal_akhir:
        query = query.lte("tanggal", tanggal_akhir)
    if q:
        query = query.ilike("keterangan", f"%{q}%")

    data = query.order("tanggal", desc=True).execute().data
    kategori_list = supabase.table("kategori").select("*").execute().data

    return render_template("transaksi/list.html",
                           transaksi=data,
                           kategori_list=kategori_list,
                           filter_jenis=jenis,
                           filter_kategori_id=kategori_id,
                           filter_tanggal_mulai=tanggal_mulai,
                           filter_tanggal_akhir=tanggal_akhir,
                           filter_q=q)


@app.route("/transaksi/tambah", methods=["GET", "POST"])
@login_required
def tambah_transaksi():
    kategori_list = supabase.table("kategori").select("*").execute().data

    if request.method == "POST":
        bukti_url = upload_bukti(request.files.get("bukti"))

        supabase.table("transaksi").insert({
            "tanggal": request.form.get("tanggal") or str(date.today()),
            "jenis": request.form["jenis"],
            "kategori_id": request.form["kategori_id"],
            "jumlah": float(request.form["jumlah"]),
            "keterangan": request.form.get("keterangan"),
            "metode": request.form.get("metode", "cash"),
            "bukti_url": bukti_url,
        }).execute()
        flash("Transaksi berhasil disimpan.", "success")
        return redirect(url_for("list_transaksi"))

    return render_template("transaksi/form.html", kategori_list=kategori_list)


@app.route("/transaksi/edit/<int:id>", methods=["GET", "POST"])
@login_required
@admin_required
def edit_transaksi(id):
    kategori_list = supabase.table("kategori").select("*").execute().data

    if request.method == "POST":
        data_update = {
            "tanggal": request.form["tanggal"],
            "jenis": request.form["jenis"],
            "kategori_id": request.form["kategori_id"],
            "jumlah": float(request.form["jumlah"]),
            "keterangan": request.form.get("keterangan"),
            "metode": request.form.get("metode", "cash"),
        }
        bukti_url_baru = upload_bukti(request.files.get("bukti"))
        if bukti_url_baru:
            data_lama = supabase.table("transaksi").select("bukti_url").eq("id", id).single().execute().data
            if data_lama:
                hapus_bukti_storage(data_lama.get("bukti_url"))
            data_update["bukti_url"] = bukti_url_baru

        supabase.table("transaksi").update(data_update).eq("id", id).execute()
        flash("Transaksi berhasil diperbarui.", "success")
        return redirect(url_for("list_transaksi"))

    data = supabase.table("transaksi").select("*").eq("id", id).single().execute().data
    return render_template("transaksi/form.html", kategori_list=kategori_list, transaksi=data)


def hapus_bukti_storage(bukti_url):
    """Hapus file bukti dari Supabase Storage berdasarkan URL yang tersimpan."""
    if not bukti_url:
        return
    nama_file = bukti_url.split("/bukti-transaksi/")[-1]
    if nama_file:
        try:
            supabase.storage.from_("bukti-transaksi").remove([nama_file])
        except Exception as e:
            print(f"Gagal hapus file storage: {e}")


@app.route("/transaksi/hapus/<int:id>", methods=["POST"])
@login_required
@admin_required
def hapus_transaksi(id):
    data = supabase.table("transaksi").select("bukti_url").eq("id", id).single().execute().data
    if data:
        hapus_bukti_storage(data.get("bukti_url"))

    supabase.table("transaksi").delete().eq("id", id).execute()
    flash("Transaksi berhasil dihapus.", "success")
    return redirect(url_for("list_transaksi"))


@app.route("/kategori", methods=["GET", "POST"])
@login_required
@admin_required
def kategori():
    if request.method == "POST":
        supabase.table("kategori").insert({
            "nama": request.form["nama"],
            "jenis": request.form["jenis"],
        }).execute()
        return redirect(url_for("kategori"))

    data = supabase.table("kategori").select("*").execute().data
    return render_template("kategori/list.html", kategori=data)


@app.route("/kategori/hapus/<int:id>", methods=["POST"])
@login_required
@admin_required
def hapus_kategori(id):
    dipakai = supabase.table("transaksi").select("id").eq("kategori_id", id).limit(1).execute().data
    if dipakai:
        flash("Kategori ini tidak bisa dihapus karena masih dipakai di transaksi yang ada.", "error")
        return redirect(url_for("kategori"))

    supabase.table("kategori").delete().eq("id", id).execute()
    flash("Kategori berhasil dihapus.", "success")
    return redirect(url_for("kategori"))

@app.route("/users", methods=["GET", "POST"])
@login_required
@admin_required
def users():
    if request.method == "POST":
        username = request.form["username"].strip()
        nama = request.form.get("nama", "").strip()
        password = request.form["password"]

        existing = supabase.table("users").select("id").eq("username", username).execute().data
        if existing:
            flash("Username sudah dipakai, pilih username lain.", "error")
            return redirect(url_for("users"))

        role = request.form.get("role", "staff")
        supabase.table("users").insert({
            "username": username,
            "nama": nama,
            "password_hash": generate_password_hash(password),
            "role": role,
        }).execute()
        flash(f"User '{username}' berhasil ditambahkan.", "success")
        return redirect(url_for("users"))

    data = supabase.table("users").select("id, username, nama, role, created_at").order("id").execute().data
    return render_template("users/list.html", users=data)


@app.route("/users/hapus/<int:id>", methods=["POST"])
@login_required
@admin_required
def hapus_user(id):
    if str(id) == str(current_user.id):
        flash("Tidak bisa menghapus akun yang sedang login.", "error")
        return redirect(url_for("users"))

    total_user = supabase.table("users").select("id").execute().data
    if len(total_user) <= 1:
        flash("Tidak bisa menghapus, minimal harus ada 1 user.", "error")
        return redirect(url_for("users"))

    supabase.table("users").delete().eq("id", id).execute()
    flash("User berhasil dihapus.", "success")
    return redirect(url_for("users"))

@app.route("/laporan")
@login_required
def laporan():
    bulan = request.args.get("bulan")
    query = supabase.table("transaksi").select("*, kategori(nama)")
    if bulan:
        tahun, bln = map(int, bulan.split("-"))
        hari_terakhir = calendar.monthrange(tahun, bln)[1]
        query = query.gte("tanggal", f"{bulan}-01").lte("tanggal", f"{bulan}-{hari_terakhir:02d}")
    transaksi = query.order("tanggal").execute().data

    rekap = {}
    for t in transaksi:
        nama_kat = t["kategori"]["nama"] if t.get("kategori") else "Tanpa Kategori"
        key = (nama_kat, t["jenis"])
        rekap[key] = rekap.get(key, 0) + t["jumlah"]

    total_masuk = sum(t["jumlah"] for t in transaksi if t["jenis"] == "masuk")
    total_keluar = sum(t["jumlah"] for t in transaksi if t["jenis"] == "keluar")

    return render_template("laporan.html",
                           rekap=rekap,
                           total_masuk=total_masuk,
                           total_keluar=total_keluar,
                           bulan=bulan)
def get_data_laporan(bulan):
    query = supabase.table("transaksi").select("*, kategori(nama)")
    if bulan:
        tahun, bln = map(int, bulan.split("-"))
        hari_terakhir = calendar.monthrange(tahun, bln)[1]
        query = query.gte("tanggal", f"{bulan}-01").lte("tanggal", f"{bulan}-{hari_terakhir:02d}")
    return query.order("tanggal").execute().data


@app.route("/laporan/export/excel")
@login_required
def export_excel():
    bulan = request.args.get("bulan")
    transaksi = get_data_laporan(bulan)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Laporan Transaksi"

    headers = ["Tanggal", "Jenis", "Kategori", "Jumlah", "Keterangan", "Metode"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F3A5F", end_color="1F3A5F", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    total_masuk = 0
    total_keluar = 0

    for t in transaksi:
        nama_kat = t["kategori"]["nama"] if t.get("kategori") else "-"
        ws.append([
            t["tanggal"],
            "Masuk" if t["jenis"] == "masuk" else "Keluar",
            nama_kat,
            t["jumlah"],
            t.get("keterangan") or "-",
            t.get("metode") or "-",
        ])
        r = ws.max_row
        amount_cell = ws.cell(row=r, column=4)
        amount_cell.number_format = "#,##0"
        if t["jenis"] == "masuk":
            amount_cell.font = Font(color="2F6F4E")
            total_masuk += t["jumlah"]
        else:
            amount_cell.font = Font(color="9C3B2A")
            total_keluar += t["jumlah"]

    ws.append([])
    total_row = ws.max_row + 1
    ws.cell(row=total_row, column=3, value="Total Masuk").font = Font(bold=True)
    ws.cell(row=total_row, column=4, value=total_masuk).font = Font(bold=True, color="2F6F4E")
    ws.cell(row=total_row, column=4).number_format = "#,##0"

    ws.cell(row=total_row + 1, column=3, value="Total Keluar").font = Font(bold=True)
    ws.cell(row=total_row + 1, column=4, value=total_keluar).font = Font(bold=True, color="9C3B2A")
    ws.cell(row=total_row + 1, column=4).number_format = "#,##0"

    ws.cell(row=total_row + 2, column=3, value="Saldo").font = Font(bold=True)
    ws.cell(row=total_row + 2, column=4, value=total_masuk - total_keluar).font = Font(bold=True)
    ws.cell(row=total_row + 2, column=4).number_format = "#,##0"

    col_widths = [14, 10, 20, 15, 30, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    nama_file = f"laporan-{bulan or 'semua'}.xlsx"
    return send_file(output, download_name=nama_file, as_attachment=True,
                      mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/laporan/export/pdf")
@login_required
def export_pdf():
    bulan = request.args.get("bulan")
    transaksi = get_data_laporan(bulan)

    total_masuk = sum(t["jumlah"] for t in transaksi if t["jenis"] == "masuk")
    total_keluar = sum(t["jumlah"] for t in transaksi if t["jenis"] == "keluar")

    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title", parent=styles["Heading1"], textColor=colors.HexColor("#1D2B22"))

    elements = []
    judul = f"Laporan Transaksi - {bulan}" if bulan else "Laporan Transaksi - Semua Periode"
    elements.append(Paragraph(judul, title_style))
    elements.append(Spacer(1, 12))

    data = [["Tanggal", "Jenis", "Kategori", "Jumlah", "Keterangan"]]
    for t in transaksi:
        nama_kat = t["kategori"]["nama"] if t.get("kategori") else "-"
        jumlah_fmt = f"Rp {t['jumlah']:,.0f}"
        data.append([
            t["tanggal"],
            "Masuk" if t["jenis"] == "masuk" else "Keluar",
            nama_kat,
            jumlah_fmt,
            t.get("keterangan") or "-",
        ])

    table = Table(data, colWidths=[70, 50, 90, 80, 140])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B7C4AC")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8F0")]),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 16))

    ringkasan_data = [
        ["Total Masuk", f"Rp {total_masuk:,.0f}"],
        ["Total Keluar", f"Rp {total_keluar:,.0f}"],
        ["Saldo", f"Rp {total_masuk - total_keluar:,.0f}"],
    ]
    ringkasan_table = Table(ringkasan_data, colWidths=[100, 100])
    ringkasan_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("TEXTCOLOR", (0, 0), (0, 0), colors.HexColor("#2F6F4E")),
        ("TEXTCOLOR", (1, 0), (1, 0), colors.HexColor("#2F6F4E")),
        ("TEXTCOLOR", (0, 1), (0, 1), colors.HexColor("#9C3B2A")),
        ("TEXTCOLOR", (1, 1), (1, 1), colors.HexColor("#9C3B2A")),
    ]))
    elements.append(ringkasan_table)

    doc.build(elements)
    output.seek(0)

    nama_file = f"laporan-{bulan or 'semua'}.pdf"
    return send_file(output, download_name=nama_file, as_attachment=True, mimetype="application/pdf")

if __name__ == "__main__":
    app.run(debug=True)